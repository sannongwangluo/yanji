# -*- coding: utf-8 -*-
"""docx 排版单测：卡片（引用块）与表格（Markdown 表格）+ 旧三档能力回归。

色值/尺寸断言全部对齐钉钉「智能纪要」样例 XML 的实测值：
卡片 = 单格表格 + 只留左侧竖线（single sz=18 BBBFC4）+ 无底纹 + 内边距 60/120/30/120
+ 文字色 646a73；表格 = Table Grid + 表头加粗 + DEE0E3 底纹。

**期望值一律写字面值**（不拿 docx_writer 的常量当期望）：常量被改坏时测试必须变红，
拿常量互证是自参照，变异实验实测会照绿。
"""
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from docx.oxml.ns import qn

from docx_writer import (_add_card, _add_markdown_table, _render_markdown,
                         _save_docx_at, save_minutes_pair)

_MD_WITH_CARD_AND_TABLE = """# 智能纪要：测试主题 2026年9月16日

> **录音主题**：测试主题
> **录音时间**：2026年9月16日（周三） 15:46
> **参会人员**：徐志龙、飞明

## 总结
音频围绕测试展开研讨，明确三件事，内容如下：

### 价格与三人团
**调价与报价线**
- 讨论64块3、13块5；六包调至14。

## 智能章节
### 00:00 三人团价格
> 会上讨论调价和三人团报名，提到64块3。

## 关键决策
| 渠道 | 价格 | 费比 |
| --- | --- | --- |
| 天猫 | 84 | 15% |
| 拼多多 | 14 | 10% |
"""


def _render(md):
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "t.docx")
        _save_docx_at(md, path)
        return Document(path)


def _content_paras(doc):
    """正文段落（跳过顶部「生成时间：…」元信息段与工具自带「会议纪要」大标题）。"""
    return [p for p in doc.paragraphs
            if p.text.strip() and not p.text.startswith("生成时间：")
            and p.text.strip() != "会议纪要"]


def _tc_pr(cell):
    return cell._tc.find(qn("w:tcPr"))


def _cell_borders(cell):
    borders = _tc_pr(cell).find(qn("w:tcBorders"))
    return {el.tag.split("}")[-1]: (el.get(qn("w:val")), el.get(qn("w:sz")),
                                    el.get(qn("w:color")))
            for el in (borders if borders is not None else [])}


def _cell_margins(cell):
    mar = _tc_pr(cell).find(qn("w:tcMar"))
    return {el.tag.split("}")[-1]: el.get(qn("w:w"))
            for el in (mar if mar is not None else [])}


class CardTest(unittest.TestCase):
    """引用块 → 卡片（样例实测形态）。"""

    def test_quote_block_becomes_single_cell_card(self):
        doc = _render("> 第一行\n> 第二行\n")
        self.assertEqual(len(doc.tables), 1)
        table = doc.tables[0]
        self.assertEqual((len(table.rows), len(table.columns)), (1, 1))
        cell = table.cell(0, 0)
        self.assertIn("第一行", cell.text)
        self.assertIn("第二行", cell.text)
        self.assertEqual(len(cell.paragraphs), 2)     # 每行一段（同样例信息卡）

    def test_card_border_is_left_bar_only(self):
        """字面值断言（不拿 dw.CARD_ACCENT 当期望值，否则常量改坏测试照绿）。"""
        doc = _render("> 内容\n")
        borders = _cell_borders(doc.tables[0].cell(0, 0))
        self.assertEqual(borders["left"], ("single", "18", "BBBFC4"))
        for side in ("top", "bottom", "right"):
            self.assertEqual(borders[side], ("nil", None, None), side)

    def test_card_has_no_shading(self):
        doc = _render("> 内容\n")
        self.assertIsNone(_tc_pr(doc.tables[0].cell(0, 0)).find(qn("w:shd")))

    def test_card_margins_match_sample(self):
        doc = _render("> 内容\n")
        self.assertEqual(_cell_margins(doc.tables[0].cell(0, 0)),
                         {"top": "60", "left": "120", "bottom": "30", "right": "120"})

    def test_card_text_color(self):
        doc = _render("> 内容一\n")
        card = doc.tables[0].cell(0, 0)
        color = card.paragraphs[0].runs[0].font.color.rgb
        self.assertEqual(str(color).upper(), "646A73")
        self.assertEqual(card.paragraphs[0].runs[0].font.size.pt, 11)

    def test_two_quote_blocks_two_cards(self):
        doc = _render("> 卡片一\n\n> 卡片二\n")
        self.assertEqual(len(doc.tables), 2)
        self.assertIn("卡片一", doc.tables[0].cell(0, 0).text)
        self.assertIn("卡片二", doc.tables[1].cell(0, 0).text)

    def test_bold_inside_card(self):
        """卡片里的 **粗体** 标签仍加粗（头部三行就是这么写的）。"""
        doc = _render("> **录音主题**：测试\n")
        runs = doc.tables[0].cell(0, 0).paragraphs[0].runs
        self.assertTrue(runs[0].bold)
        self.assertIn("录音主题", runs[0].text)


class TableTest(unittest.TestCase):
    """Markdown 表格 → 真 docx 表格。"""

    def test_table_shape_and_style(self):
        doc = _render(_MD_WITH_CARD_AND_TABLE)
        table = doc.tables[-1]
        self.assertEqual((len(table.rows), len(table.columns)), (3, 3))
        self.assertEqual(table.style.name, "Table Grid")   # 细边框

    def test_header_bold_and_shaded(self):
        doc = _render(_MD_WITH_CARD_AND_TABLE)
        table = doc.tables[-1]
        for j in range(3):
            cell = table.cell(0, j)
            self.assertTrue(cell.paragraphs[0].runs[0].bold, f"表头第 {j} 列没加粗")
            shd = _tc_pr(cell).find(qn("w:shd"))
            self.assertIsNotNone(shd, "表头没底纹")
            self.assertEqual(shd.get(qn("w:fill")), "DEE0E3")

    def test_body_cells_content(self):
        doc = _render(_MD_WITH_CARD_AND_TABLE)
        table = doc.tables[-1]
        self.assertEqual([table.cell(1, j).text for j in range(3)],
                         ["天猫", "84", "15%"])
        self.assertEqual([table.cell(2, j).text for j in range(3)],
                         ["拼多多", "14", "10%"])
        self.assertFalse(table.cell(1, 0).paragraphs[0].runs[0].bold)
        # 分隔行 | --- | 不能变成数据行
        self.assertEqual(table.cell(1, 0).text, "天猫")

    def test_ragged_row_padded(self):
        doc = _render("| A | B |\n| --- | --- |\n| 只有一个 |\n")
        table = doc.tables[0]
        self.assertEqual(table.cell(1, 1).text, "")

    def test_separator_only_table_ignored(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            self.assertIsNone(_add_markdown_table(doc, [["---", "---"]]))
            self.assertEqual(len(doc.tables), 0)


class TableEdgeCaseTest(unittest.TestCase):
    """C2-3：GFM 表格边界——行尾竖线可省 / 分隔行允许单个短横 / 转义与行内代码里的竖线。"""

    def test_row_without_trailing_pipe_is_a_table(self):
        """行尾竖线可省（GFM 合法），整张表仍要被识别成表格。"""
        doc = _render("| 渠道 | 价格\n| --- | ---\n| 天猫 | 84\n")
        self.assertEqual(len(doc.tables), 1)
        table = doc.tables[0]
        self.assertEqual((len(table.rows), len(table.columns)), (2, 2))
        self.assertEqual([table.cell(0, j).text for j in range(2)], ["渠道", "价格"])
        self.assertEqual([table.cell(1, j).text for j in range(2)], ["天猫", "84"])

    def test_separator_allows_single_dash(self):
        """分隔行允许 1 个短横（`| - | - |`），不能被当成数据行吃掉。"""
        doc = _render("| 渠道 | 价格 |\n| - | - |\n| 天猫 | 84 |\n")
        table = doc.tables[0]
        self.assertEqual((len(table.rows), len(table.columns)), (2, 2))
        self.assertEqual(table.cell(1, 0).text, "天猫")

    def test_escaped_pipe_is_literal_content(self):
        r"""未转义的 `|` 才切分，`\|` 还原成字面竖线。"""
        doc = _render("| 项 | 值 |\n| --- | --- |\n| a \\| b | c |\n")
        table = doc.tables[0]
        self.assertEqual(table.cell(1, 0).text, "a | b")
        self.assertEqual(table.cell(1, 1).text, "c")

    def test_pipe_inside_inline_code_is_not_a_separator(self):
        """行内代码 `a|b` 里的竖线不当分隔符。"""
        doc = _render("| 项 | 值 |\n| --- | --- |\n| `x|y` | c |\n")
        table = doc.tables[0]
        self.assertEqual(table.cell(1, 0).text, "x|y")
        self.assertEqual(table.cell(1, 1).text, "c")


class LegacyRenderingTest(unittest.TestCase):
    """旧三档能力不回归 + 新字号/字体。"""

    def test_heading_sizes_black_bold(self):
        doc = _render("# 大标题\n## 栏目\n### 小节\n#### 更小\n")
        paras = _content_paras(doc)
        sizes = [(p.runs[0].font.size.pt, p.runs[0].bold) for p in paras]
        self.assertEqual(sizes, [(26, True), (18, True), (12, True), (11, True)])
        for p in paras:
            self.assertIsNone(p.runs[0].font.color.rgb)   # 黑色，不是默认蓝

    def test_paragraph_spacing_literal(self):
        """正文 1.2 倍行距、段前段后 6pt；栏目标题（##）19/7pt——全部写字面值。"""
        doc = _render("正文\n\n## 栏目\n")
        body, heading = _content_paras(doc)
        self.assertEqual(body.paragraph_format.line_spacing, 1.2)
        self.assertEqual(body.paragraph_format.space_before.pt, 6)
        self.assertEqual(body.paragraph_format.space_after.pt, 6)
        self.assertEqual(heading.paragraph_format.space_before.pt, 19)
        self.assertEqual(heading.paragraph_format.space_after.pt, 7)

    def test_body_and_lists(self):
        doc = _render("正文一段\n\n- 无序项\n- 又一项\n\n1. 有序项\n")
        paras = _content_paras(doc)
        self.assertEqual(paras[0].text, "正文一段")
        self.assertEqual(paras[0].runs[0].font.size.pt, 11)
        self.assertEqual(paras[1].style.name, "List Bullet")
        self.assertEqual(paras[2].style.name, "List Bullet")
        self.assertEqual(paras[3].style.name, "List Number")
        self.assertEqual(paras[3].text, "有序项")

    def test_inline_bold_toggle(self):
        doc = _render("**加粗**普通**又加粗**\n")
        runs = _content_paras(doc)[0].runs
        self.assertEqual([(r.text, r.bold) for r in runs],
                         [("加粗", True), ("普通", False), ("又加粗", True)])

    def test_inline_bold_not_inverted_when_line_starts_bold(self):
        """回归：行首就是 ** 时不能整行加粗/普通反过来（旧实现有这个问题）。"""
        doc = _render("**关键决策**：三人团少报库存\n")
        runs = _content_paras(doc)[0].runs
        self.assertEqual([(r.text, r.bold) for r in runs],
                         [("关键决策", True), ("：三人团少报库存", False)])

    def test_fonts_are_dengxian_arial(self):
        doc = _render("正文\n")
        rPr = doc.paragraphs[0].runs[0]._element.find(qn("w:rPr"))
        rFonts = rPr.find(qn("w:rFonts"))
        self.assertEqual(rFonts.get(qn("w:eastAsia")), "等线")
        self.assertEqual(rFonts.get(qn("w:ascii")), "Arial")

    def test_no_extra_title_when_md_has_h1(self):
        """md 自带 # 大标题时，doc.paragraphs 原始段落里不得出现独立「会议纪要」标题段。

        注意断言要落在**原始** doc.paragraphs 上：helper _content_paras 已经把
        「会议纪要」滤掉了，拿它做 NotIn 是恒真的空断言。
        """
        doc = _render("# 智能纪要：主题 2026年9月16日\n正文\n")
        texts = [p.text.strip() for p in doc.paragraphs]
        self.assertNotIn("会议纪要", texts)
        content = [t for t in texts if t and not t.startswith("生成时间：")]
        self.assertTrue(content[0].startswith("智能纪要："))

    def test_keep_title_when_md_has_no_h1(self):
        doc = _render("正文\n")
        texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        self.assertEqual(texts[0], "会议纪要")   # 无 h1 时保留工具自己的大标题

    def test_full_document_cards_and_table_counts(self):
        doc = _render(_MD_WITH_CARD_AND_TABLE)
        self.assertEqual(len(doc.tables), 3)   # 头部卡 + 章节摘要卡 + 数据表

    def test_pair_export_still_works(self):
        with tempfile.TemporaryDirectory() as tmp:
            docx_path, md_path = save_minutes_pair(_MD_WITH_CARD_AND_TABLE, tmp)
            self.assertTrue(os.path.exists(docx_path))
            with open(md_path, encoding="utf-8") as f:
                self.assertEqual(f.read(), _MD_WITH_CARD_AND_TABLE)


class ControlCharacterTest(unittest.TestCase):
    """C2-1：md 里混入 XML 非法控制字符时，docx 仍必须出得来。

    实测 python-docx 的 add_run 遇 \\x00/\\x07/\\x1f 抛
    ValueError("All strings must be XML compatible…")，渲染链没有兜底 →
    md 已落盘、docx 完全出不来。渲染层统一滤掉（md 文件内容原样不动）。

    注意 \\x0b/\\x0c 走不到 add_run：`_render_markdown` 先做 `md.splitlines()`，
    而 Python 的 splitlines 把 \\x0b/\\x0c/\\x1c-\\x1e 也当行分隔（会被切成两行，
    不炸但结构变形）；真正会炸的是 \\x00-\\x08、\\x0e-\\x1b、\\x1f。两类都覆盖。
    """

    _MD = ("# 大标题\x0b带控制符\n\n"
           "**粗体\x0b部分**：正文\x0c内容\x07\n\n"
           "| 表头\x00A | B |\n| --- | --- |\n| 单元\x0b格 | 二 |\n\n"
           "> 卡片\x0b文字\n")

    _HARD_MD = ("# 标题\x00带控制符\x07\n\n"
                "**粗体\x07部分**：正文\x1f内容\n\n"
                "| 表头\x00A | B |\n| --- | --- |\n| 单元\x00格 | 二 |\n\n"
                "> 卡片\x01文字\n")

    def _assert_no_illegal_xml_chars(self, doc):
        self.assertNotRegex(doc.element.xml, r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

    def test_docx_renders_with_xml_illegal_chars(self):
        """\\x00/\\x07/\\x1f 会直接炸 add_run：修不好这里就是 ValueError。"""
        doc = _render(self._HARD_MD)
        self._assert_no_illegal_xml_chars(doc)

    def test_heading_and_body_control_chars_stripped(self):
        doc = _render(self._HARD_MD)
        texts = [p.text for p in doc.paragraphs]
        self.assertIn("标题带控制符", texts)
        self.assertIn("粗体部分：正文内容", texts)
        for text in texts:
            self.assertNotRegex(text, r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

    def test_table_cell_control_chars_stripped(self):
        doc = _render(self._HARD_MD)
        table = doc.tables[0]
        self.assertEqual([table.cell(0, j).text for j in range(2)], ["表头A", "B"])
        self.assertEqual([table.cell(1, j).text for j in range(2)], ["单元格", "二"])

    def test_card_control_chars_stripped(self):
        doc = _render(self._HARD_MD)
        cell = doc.tables[1].cell(0, 0)
        self.assertEqual(cell.text, "卡片文字")
        self.assertEqual(cell.paragraphs[0].runs[0].text, "卡片文字")

    def test_bold_parity_kept_after_filtering(self):
        doc = _render("**粗体\x07**：尾部\n")
        runs = _content_paras(doc)[0].runs
        self.assertEqual([(r.text, r.bold) for r in runs],
                         [("粗体", True), ("：尾部", False)])

    def test_vtab_and_formfeed_md_still_renders_docx(self):
        """题目口径：标题/正文/粗体/表格单元格里带 \\x0b 都要能正常出 docx。"""
        doc = _render(self._MD)
        self.assertTrue(doc.paragraphs)
        self.assertEqual(len(doc.tables), 2)     # 数据表 + 卡片
        self._assert_no_illegal_xml_chars(doc)

    def test_md_file_is_written_verbatim(self):
        """红线：md 只是文本，控制字符照原样写盘（过滤只发生在 docx 渲染层）。"""
        with tempfile.TemporaryDirectory() as tmp:
            _, md_path = save_minutes_pair(self._MD, tmp)
            with open(md_path, "rb") as f:
                self.assertEqual(f.read(), self._MD.encode("utf-8"))


class ManualDocxShadingTest(unittest.TestCase):
    """C2-10：gen_manual_docx 代码块底纹的 OOXML 顺序（w:shd 在 w:spacing/w:ind 之前）。"""

    def test_code_block_shading_order(self):
        import gen_manual_docx as gmd
        with tempfile.TemporaryDirectory() as tmp:
            md_path = os.path.join(tmp, "m.md")
            out_path = os.path.join(tmp, "m.docx")
            with open(md_path, "w", encoding="utf-8") as f:
                f.write("# 标题\n\n```\ncode line\n```\n")
            gmd.render(md_path, out_path)
            doc = Document(out_path)
            shaded = [p for p in doc.paragraphs
                      if p._element.find(qn("w:pPr")) is not None
                      and p._element.find(qn("w:pPr")).find(qn("w:shd")) is not None]
            self.assertEqual(len(shaded), 1)
            pPr = shaded[0]._element.find(qn("w:pPr"))
            tags = [el.tag.split("}")[-1] for el in pPr]
            self.assertLess(tags.index("shd"), tags.index("spacing"))
            self.assertLess(tags.index("shd"), tags.index("ind"))
            self.assertEqual(pPr.find(qn("w:shd")).get(qn("w:fill")), "F2F2F2")


if __name__ == "__main__":
    unittest.main(verbosity=2)
